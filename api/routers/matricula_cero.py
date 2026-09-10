import io
from collections import Counter
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from sqlalchemy import bindparam, text

from ..core.database import engine_analitica, engine_convocatoria
from ..core.matricula_cero_helpers import calcular_periodo_label
from ..core.mc_final_mapping import (
    COL_DOCPERIODO_ULTIMO_MC,
    COLUMNAS_ESPERADAS_MC_FINAL,
    FilaInvalidaError,
    MC_FINAL_COLUMNS,
    PLACEHOLDER_NO_ENCONTRADO,
    UPDATE_COLUMNS,
    map_row,
    normalize_text,
    to_text,
)
from ..models.consulta import ConsultaResponse
from ..models.matricula_cero import (
    CargaMcFinalResponse,
    FilaProblemaMcFinal,
    InfoPersonalMCResponse,
    ResumenCargaMcFinal,
    UltimaActualizacionTableroResponse,
)
from .seguimiento_auth import get_current_user_seguimiento, require_rol

router = APIRouter(prefix="/matricula-cero", tags=["Matrícula Cero"])

# NOTA: estos endpoints se autentican con el login de Seguimiento
# (usuarios_seg_proceso_mc), NO con el login genérico de la API — decisión
# tomada para centralizar roles bajo un solo sistema de usuarios, ya que
# quien necesita Consulta+Tablero es el mismo equipo que ya tiene (o puede
# tener) cuenta de Seguimiento. Cualquier rol autenticado (ADMIN, DIRECTORA,
# LMC, AST, AD, AF, AJ) puede consultar — es de solo lectura, no se restringe
# por rol específico aquí.


@router.get(
    "/consulta",
    response_model=ConsultaResponse,
    summary="Consultar formulario de Matrícula Cero (vista vigente 2026-2)",
)
def consulta_formulario_2026_2(
    documento: str = Query(..., min_length=6, max_length=15),
    _: Dict[str, Any] = Depends(get_current_user_seguimiento),
) -> ConsultaResponse:
    """
    Endpoint NUEVO y separado de /consulta/formulario-mc (que se deja intacto,
    apuntando a la vista 2025-2). Este consulta la vista vigente
    vw_matricula_cero_2026_2 — equivalente a
    app/database/db_operations.py::MatriculaCeroOperations.get_by_documento
    del portal Streamlit.
    """
    q = text("""
        SELECT *
        FROM vw_matricula_cero_2026_2
        WHERE documento = :documento
        ORDER BY fecha_registro DESC
    """)
    with engine_convocatoria.connect() as conn:
        rows = conn.execute(q, {"documento": documento}).mappings().all()

    results: List[Dict[str, Any]] = [dict(r) for r in rows]
    return ConsultaResponse(count=len(results), results=results)


@router.get(
    "/tablero/info-personal",
    response_model=InfoPersonalMCResponse,
    summary="Información personal del formulario más reciente (Tablero Matrícula Cero)",
)
def tablero_info_personal(
    documento: str = Query(..., min_length=6, max_length=15),
    _: Dict[str, Any] = Depends(get_current_user_seguimiento),
) -> InfoPersonalMCResponse:
    """
    Equivalente a tablero_mc.py::_cargar_info_personal: consulta directamente
    convocatoria_sapiencia.matricula_cero (con sus tablas de referencia/catálogo)
    filtrando el ÚLTIMO período registrado para el documento — sin depender de
    ninguna vista. A diferencia del original (que arma la consulta con
    f-strings), aquí el documento va parametrizado para evitar inyección SQL.
    """
    q = text("""
        SELECT mc.*,
               a.nombre  AS tipo_documento,
               b.nombre  AS pais_nacimiento,
               d.nombre  AS departamento_nacimiento,
               c.nombre  AS municipio_nacimiento,
               e.nombre  AS sexo,
               f.nombre  AS orientacion_sexual,
               g.nombre  AS identidad_genero,
               i.nombre  AS afiliacion_salud,
               j.nombre  AS tipo_vivienda,
               k.nombre  AS actividad_realiza,
               m.nombre  AS estrato,
               n.nombre  AS pais_residencia_ubg,
               o.nombre  AS departamento_ubicacion,
               p.nombre  AS municipio_ubicacion,
               q.nombre  AS barrio,
               r.nombre  AS comuna,
               w.nombre  AS nivel_academico,
               yy.nombre AS beneficio_sapiencia,
               z.nombre  AS ies_adscritas,
               zz.nombre AS programa_admitido,
               aa.nombre AS semestre_academico,
               'Completo' AS estado_formulario
        FROM convocatoria_sapiencia.matricula_cero mc
        LEFT JOIN convocatoria_sapiencia.vlf_tipo_documento      a  ON a.id  = mc.tipo_documento
        LEFT JOIN convocatoria_sapiencia.pais                    b  ON b.id  = mc.pais_nacimiento
        LEFT JOIN convocatoria_sapiencia.departamento            d  ON d.id  = mc.departamento_residencia
        LEFT JOIN convocatoria_sapiencia.municipio               c  ON c.id  = mc.municipio_residencia
        LEFT JOIN convocatoria_sapiencia.odes_expectativas_sexo  e  ON e.id  = mc.sexo
        LEFT JOIN convocatoria_sapiencia.estudiantes_orientacion_sexual f ON f.id = mc.orientacion_sexual
        LEFT JOIN convocatoria_sapiencia.estudiantes_identidad_genero   g ON g.id = mc.identidad_genero
        LEFT JOIN convocatoria_sapiencia.tipo_regimen_salud      i  ON i.id  = mc.afiliacion_salud
        LEFT JOIN convocatoria_sapiencia.talento_especializado_tipo_vivienda j ON j.id = mc.tipo_vivienda
        LEFT JOIN convocatoria_sapiencia.actividad_matricula_cero k  ON k.id  = mc.actividad_realiza
        LEFT JOIN convocatoria_sapiencia.vlf_estrato             m  ON m.id  = mc.estrato
        LEFT JOIN convocatoria_sapiencia.pais                    n  ON n.id  = mc.pais_residencia_ubg
        LEFT JOIN convocatoria_sapiencia.departamento            o  ON o.id  = mc.departamento_ubg
        LEFT JOIN convocatoria_sapiencia.municipio               p  ON p.id  = mc.municipio_residencia_ubg
        LEFT JOIN convocatoria_sapiencia.barrio                  q  ON q.id  = mc.barrio
        LEFT JOIN convocatoria_sapiencia.comuna_caracterizacion  r  ON r.id  = mc.comuna
        LEFT JOIN convocatoria_sapiencia.nivel_academico_matricula_cero w ON w.id = mc.nivel_academico
        LEFT JOIN convocatoria_sapiencia.matriculacero_beneficio_sapiencia yy ON yy.id = mc.beneficio_sapiencia
        LEFT JOIN convocatoria_sapiencia.ies_acoso_sexual        z  ON z.id  = mc.ies_adscritas
        LEFT JOIN convocatoria_sapiencia.ies_matricula_cero_actual zz ON zz.id = mc.programa_admitido
        LEFT JOIN convocatoria_sapiencia.semestre_matricula_cero aa ON aa.id = mc.semestre_academico
        WHERE mc.documento = :documento
          AND mc.periodo = (
              SELECT MAX(mc2.periodo)
              FROM convocatoria_sapiencia.matricula_cero mc2
              WHERE mc2.documento = :documento
          )
        LIMIT 1
    """)
    with engine_convocatoria.connect() as conn:
        row = conn.execute(q, {"documento": documento}).mappings().fetchone()

    if not row:
        return InfoPersonalMCResponse(encontrado=False)

    datos = dict(row)
    periodo_label = calcular_periodo_label(datos.get("periodo"))
    return InfoPersonalMCResponse(encontrado=True, periodo_label=periodo_label, datos=datos)


@router.get(
    "/tablero/giros",
    response_model=ConsultaResponse,
    summary="Historial de giros/seguimiento académico por período (Tablero Matrícula Cero)",
)
def tablero_giros(
    documento: str = Query(..., min_length=6, max_length=15),
    solo_proyecto: bool = Query(
        True,
        description="Si es True (por defecto), solo incluye períodos desde 2023-2 en adelante "
                    "(el proyecto actual). Si es False, incluye todo el histórico.",
    ),
    _: Dict[str, Any] = Depends(get_current_user_seguimiento),
) -> ConsultaResponse:
    """
    Equivalente a tablero_mc.py::_cargar_analitica + el filtro
    df_analitica_filtrado (columna 'periodo' es texto tipo 'AAAA-S' en esta
    tabla — distinto del código entero de matricula_cero — se compara igual
    que en el original, lexicográficamente).
    """
    q = text("""
        SELECT *
        FROM analitica_fondos.mc_final
        WHERE documento = :documento
        ORDER BY periodo ASC
    """)
    with engine_analitica.connect() as conn:
        rows = conn.execute(q, {"documento": documento}).mappings().all()

    results: List[Dict[str, Any]] = [dict(r) for r in rows]
    if solo_proyecto:
        results = [r for r in results if str(r.get("periodo", "")) >= "2023-2"]

    return ConsultaResponse(count=len(results), results=results)


@router.get(
    "/tablero/ultima-actualizacion",
    response_model=UltimaActualizacionTableroResponse,
    summary="Fecha del cargue más reciente de datos históricos en el Tablero",
)
def ultima_actualizacion_tablero(
    _: Dict[str, Any] = Depends(get_current_user_seguimiento),
) -> UltimaActualizacionTableroResponse:
    """A pedido de Migue: tarjeta de "fecha de actualización" para el
    Tablero histórico. NO aplica a /consulta ni a /tablero/info-personal
    (esas consultan convocatoria_sapiencia en tiempo real, sin cargue de
    por medio) — solo al historial de mc_final, que sí se carga por
    lotes."""
    with engine_analitica.connect() as conn:
        fila = conn.execute(text("SELECT MAX(fecha_cargue) AS fecha FROM analitica_fondos.mc_final")).mappings().fetchone()
    return UltimaActualizacionTableroResponse(fecha=fila["fecha"] if fila else None)


# ---------------------------------------------------------------------------
# POST /matricula-cero/tablero/cargar-mc-final — a pedido de Migue: portar acá
# el flujo de carga que hoy corre por fuera con un script de Python aparte
# (mapping.py/main.py/db.py). Ver el docstring largo en
# api/models/matricula_cero.py (CargaMcFinalResponse y compañía) para el
# diseño completo de las categorías de "problema" — se decidió con Migue en
# conversación, no es una elección arbitraria.
#
# Diferencias clave frente a POST /reportes-financieros/cargar-excel:
#   - Un archivo por IES (Migue las recibe así, no las 8 juntas) — no hace
#     falta un selector de múltiples archivos.
#   - Archivos MUCHO más densos (hasta ~24 mil filas en IES grandes como
#     ITM) — la previsualización NO lista fila por fila (inmanejable a ese
#     tamaño): solo un resumen con conteos + el detalle de las filas que sí
#     necesitan revisión antes de confirmar.
#   - mc_final no tiene UNIQUE KEY en docperiodo (duplicados históricos
#     pendientes de limpiar a mano — ver comentario en la migración), así
#     que no se puede usar "INSERT ... ON DUPLICATE KEY UPDATE" como en
#     Financiero. En su lugar, se clasifica cada fila ANTES de escribir:
#     0 registros existentes -> insertar; 1 -> actualizar (si la IES
#     coincide) o "conflicto_ies" (si no); 2+ -> "docperiodo_ambiguo", no se
#     toca. Mismo criterio que usaba el script aparte de Migue.
# ---------------------------------------------------------------------------

_MAX_BYTES_MC_FINAL = 30 * 1024 * 1024  # 30 MB — archivos más densos que Financiero
_MAX_FILAS_MC_FINAL = 40000  # ITM (la IES más grande) ronda 24k; margen de sobra
_CHUNK_SIZE = 1000  # tamaño seguro de cada IN (...) / executemany, igual que el script aparte
_SHEET_DEFAULT = "ANALISIS"
_SUBIR_COL_DEFAULT = "SUBIR_A_BD"
_CAMPOS_RENUEVA_PLACEHOLDER = ("semestre_ingreso", "giros_proyectados_totales", "giros_realizados")


def _chunked(seq, size=_CHUNK_SIZE):
    seq = list(seq)
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


@router.post(
    "/tablero/cargar-mc-final",
    response_model=CargaMcFinalResponse,
    summary="Cargar el excel de conciliación de una IES hacia mc_final (previsualizar o aplicar)",
)
async def cargar_mc_final(
    archivo: UploadFile = File(..., description="Excel de conciliación de UNA IES (hoja 'ANALISIS' por defecto)"),
    periodo: str = Form(..., description='Período de esta carga, ej. "2026-1".'),
    estado_col: str = Form(..., description='Nombre literal de la columna de estado semestral, ej. "ESTADO_(RESULTADO_26-1)".'),
    motivo_col: str = Form(..., description='Nombre literal de la columna de motivo del estado, ej. "MOTIVO_ESTADO_26-1".'),
    sheet: str = Form(_SHEET_DEFAULT, description=f'Nombre de la hoja con los datos (default: "{_SHEET_DEFAULT}").'),
    subir_col: str = Form(
        _SUBIR_COL_DEFAULT,
        description='Columna de validación (SI/NO) para excluir filas ya identificadas manualmente '
                    f'(default: "{_SUBIR_COL_DEFAULT}"). Si el archivo no la tiene, no se excluye nada.',
    ),
    fecha_cargue: Optional[date] = Form(None, description="Fecha a guardar en fecha_cargue. Por defecto, hoy."),
    confirmar: bool = Form(False, description="false = solo valida y previsualiza; true = aplica los cambios"),
    _: Dict[str, Any] = Depends(require_rol("ADMIN", "AD")),
) -> CargaMcFinalResponse:
    """
    Se llama 2 veces desde el frontend con el MISMO archivo (mismo patrón que
    Financiero): primero con confirmar=false (solo clasifica, no escribe
    nada) y, si Migue confirma, otra vez con confirmar=true (recién ahí
    aplica los INSERT/UPDATE). La clasificación se calcula IGUAL en ambas
    llamadas (contra el estado de la BD al momento de cada llamada), así que
    lo que se ve en la previsualización es lo que se aplica al confirmar —
    salvo que la BD haya cambiado entre medio (ej. alguien más cargó otra
    IES justo en ese momento).
    """
    contenido = await archivo.read()
    if len(contenido) > _MAX_BYTES_MC_FINAL:
        raise HTTPException(status_code=400, detail="El archivo supera el tamaño máximo permitido (30 MB).")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        if sheet not in wb.sheetnames:
            raise HTTPException(status_code=400, detail=f"El archivo no tiene una hoja llamada '{sheet}'.")
        ws = wb[sheet]
        filas_iter = ws.iter_rows()
        primera_fila = next(filas_iter)
    except HTTPException:
        raise
    except StopIteration:
        raise HTTPException(status_code=400, detail="El archivo está vacío.")
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo — ¿es un .xlsx válido?")

    encabezados = [str(c.value).strip() if c.value is not None else None for c in primera_fila]
    idx = {h: i for i, h in enumerate(encabezados) if h}

    # Columnas requeridas: las que siempre debe traer el excel de
    # conciliación (COLUMNAS_ESPERADAS_MC_FINAL, ver el porqué de las 2
    # excepciones — SUBIR_A_BD y NETO_PAGAR_AJUSTE_1.5 — en
    # mc_final_mapping.py) + las 2 dinámicas de este período puntual. A
    # pedido de Migue: si falta cualquiera, se rechaza el archivo completo
    # con un error que nombra TODAS las que faltan, en vez de procesar esa
    # columna en blanco silenciosamente.
    columnas_requeridas = list(dict.fromkeys(COLUMNAS_ESPERADAS_MC_FINAL + [estado_col, motivo_col]))
    faltantes = [c for c in columnas_requeridas if c not in idx]
    if faltantes:
        raise HTTPException(
            status_code=400,
            detail=f"El archivo no tiene el formato esperado — faltan las columnas: {', '.join(faltantes)}.",
        )

    def fila_a_dict(fila) -> Dict[str, Any]:
        return {header: (fila[i].value if i < len(fila) else None) for header, i in idx.items()}

    filas_crudas: List[Tuple[int, Dict[str, Any]]] = []
    numero_fila = 1
    for fila in filas_iter:
        numero_fila += 1
        if all(c.value is None for c in fila):
            continue
        filas_crudas.append((numero_fila, fila_a_dict(fila)))

    if not filas_crudas:
        raise HTTPException(status_code=400, detail="El archivo no tiene filas de datos (solo el encabezado).")
    if len(filas_crudas) > _MAX_FILAS_MC_FINAL:
        raise HTTPException(status_code=400, detail=f"El archivo tiene demasiadas filas (máximo {_MAX_FILAS_MC_FINAL}).")

    fecha_cargue_final = fecha_cargue or date.today()

    # --- 1. Excluir manualmente por SUBIR_A_BD=NO (si la columna existe) ---
    filas_excluidas: List[FilaProblemaMcFinal] = []
    filas_a_procesar: List[Tuple[int, Dict[str, Any]]] = []
    for numero_fila, row in filas_crudas:
        valor_subir = row.get(subir_col) if subir_col in idx else None
        valor_norm = "" if valor_subir is None else str(valor_subir).strip().upper()
        if subir_col in idx and valor_norm == "NO":
            filas_excluidas.append(FilaProblemaMcFinal(
                fila_excel=numero_fila,
                documento=to_text(row.get("DOCUMENTO")),
                ies=to_text(row.get("IES")),
                tipo="excluida_manual",
                mensaje=f"Excluida manualmente por columna '{subir_col}' = NO.",
            ))
        else:
            filas_a_procesar.append((numero_fila, row))

    # --- 2. Bulk lookup del histórico de RENUEVA (docperiodo del último
    #        período en MC que ya trae calculado el propio archivo) ---
    needed_historial = set()
    for _num, row in filas_a_procesar:
        if to_text(row.get("NUEVO_RENUEVA")) == "RENUEVA":
            dp_hist = to_text(row.get(COL_DOCPERIODO_ULTIMO_MC))
            if dp_hist:
                needed_historial.add(dp_hist)

    historial: Dict[str, Optional[Dict[str, Any]]] = {}
    with engine_analitica.connect() as conn:
        if needed_historial:
            stmt_hist = text("""
                SELECT docperiodo, semestre_ingreso, giros_proyectados_totales, giros_realizados
                FROM analitica_fondos.mc_final WHERE docperiodo IN :dps
            """).bindparams(bindparam("dps", expanding=True))
            agrupado: Dict[str, List[Dict[str, Any]]] = {}
            for chunk in _chunked(needed_historial):
                for r in conn.execute(stmt_hist, {"dps": tuple(chunk)}).mappings().all():
                    agrupado.setdefault(r["docperiodo"], []).append(dict(r))
            for dp_hist, coincidencias in agrupado.items():
                # 2+ coincidencias históricas ambiguas se tratan igual que
                # "no encontrado" (mismo criterio que el script aparte) — no
                # hay forma de saber cuál de las 2 es la correcta.
                historial[dp_hist] = coincidencias[0] if len(coincidencias) == 1 else None

        def lookup_previous_mc(dp_hist: str) -> Optional[Dict[str, Any]]:
            return historial.get(dp_hist)

        # --- 3. Mapear cada fila (map_row ya valida DOCUMENTO) ---
        filas_invalidas: List[FilaProblemaMcFinal] = []
        filas_mapeadas: List[Tuple[int, Dict[str, Any]]] = []
        for numero_fila, row in filas_a_procesar:
            try:
                mapeada = map_row(
                    row, periodo=periodo, estado_col=estado_col, motivo_col=motivo_col,
                    fecha_cargue=fecha_cargue_final, db_lookup_previous_mc=lookup_previous_mc,
                )
                filas_mapeadas.append((numero_fila, mapeada))
            except FilaInvalidaError as e:
                filas_invalidas.append(FilaProblemaMcFinal(
                    fila_excel=numero_fila,
                    documento=None,
                    ies=to_text(row.get("IES")),
                    tipo="invalida",
                    mensaje=str(e),
                ))

        # --- 4. RENUEVA sin historial (informativo — la fila se sube igual,
        #        pero Migue quiere verla para revisar el dato incompleto) ---
        filas_renueva_sin_historial: List[FilaProblemaMcFinal] = []
        for numero_fila, mapeada in filas_mapeadas:
            if any(mapeada.get(c) == PLACEHOLDER_NO_ENCONTRADO for c in _CAMPOS_RENUEVA_PLACEHOLDER):
                filas_renueva_sin_historial.append(FilaProblemaMcFinal(
                    fila_excel=numero_fila,
                    documento=mapeada["documento"],
                    ies=mapeada["ies"],
                    tipo="renueva_sin_historial",
                    mensaje="RENUEVA: no se encontró (o hay duplicados ambiguos para) su período anterior en "
                            "mc_final — semestre_ingreso/giros quedaron en 'NO ENCONTRADO'.",
                ))

        # --- 5. Duplicado DENTRO del mismo archivo (mismo docperiodo 2+
        #        veces) — a pedido de Migue: ninguna de esas filas se aplica,
        #        se reportan todas para que corrija el archivo. ---
        conteo_en_archivo = Counter(m["docperiodo"] for _, m in filas_mapeadas)
        docperiodos_duplicados = {dp for dp, n in conteo_en_archivo.items() if n > 1}

        filas_duplicadas_en_archivo: List[FilaProblemaMcFinal] = []
        filas_candidatas: List[Tuple[int, Dict[str, Any]]] = []
        for numero_fila, mapeada in filas_mapeadas:
            if mapeada["docperiodo"] in docperiodos_duplicados:
                filas_duplicadas_en_archivo.append(FilaProblemaMcFinal(
                    fila_excel=numero_fila,
                    documento=mapeada["documento"],
                    ies=mapeada["ies"],
                    tipo="duplicado_en_archivo",
                    mensaje=f"El documento+período aparece {conteo_en_archivo[mapeada['docperiodo']]} veces en "
                            "este mismo archivo — ninguna de esas filas se aplicó, corrige el archivo y vuelve a "
                            "subirlo.",
                ))
            else:
                filas_candidatas.append((numero_fila, mapeada))

        # --- 6. Contra la BD: 0 existentes -> insertar; 1 -> actualizar (si
        #        la IES coincide) o conflicto de IES; 2+ -> ambiguo. ---
        docperiodos_candidatos = [m["docperiodo"] for _, m in filas_candidatas]
        existing_counts: Dict[str, int] = {}
        existing_ies: Dict[str, str] = {}
        for chunk in _chunked(docperiodos_candidatos):
            if not chunk:
                continue
            stmt_cnt = text("""
                SELECT docperiodo, COUNT(*) AS n FROM analitica_fondos.mc_final
                WHERE docperiodo IN :dps GROUP BY docperiodo
            """).bindparams(bindparam("dps", expanding=True))
            for r in conn.execute(stmt_cnt, {"dps": tuple(chunk)}).mappings().all():
                existing_counts[r["docperiodo"]] = r["n"]

        con_una_coincidencia = [dp for dp in docperiodos_candidatos if existing_counts.get(dp, 0) == 1]
        for chunk in _chunked(con_una_coincidencia):
            if not chunk:
                continue
            stmt_ies = text(
                "SELECT docperiodo, ies FROM analitica_fondos.mc_final WHERE docperiodo IN :dps"
            ).bindparams(bindparam("dps", expanding=True))
            for r in conn.execute(stmt_ies, {"dps": tuple(chunk)}).mappings().all():
                existing_ies[r["docperiodo"]] = r["ies"]

        filas_conflicto_ies: List[FilaProblemaMcFinal] = []
        filas_docperiodo_ambiguo: List[FilaProblemaMcFinal] = []
        filas_para_aplicar: List[Tuple[int, Dict[str, Any], str]] = []
        for numero_fila, mapeada in filas_candidatas:
            dp = mapeada["docperiodo"]
            n_existentes = existing_counts.get(dp, 0)
            if n_existentes == 0:
                filas_para_aplicar.append((numero_fila, mapeada, "insertar"))
            elif n_existentes == 1:
                ies_actual = existing_ies.get(dp)
                if normalize_text(ies_actual) == normalize_text(mapeada["ies"]):
                    filas_para_aplicar.append((numero_fila, mapeada, "actualizar"))
                else:
                    filas_conflicto_ies.append(FilaProblemaMcFinal(
                        fila_excel=numero_fila,
                        documento=mapeada["documento"],
                        ies=mapeada["ies"],
                        tipo="conflicto_ies",
                        mensaje=f"Ya existe en mc_final bajo la IES '{ies_actual}'; este archivo trae "
                                f"'{mapeada['ies']}'. No se sobrescribió — decide manualmente cuál IES es la "
                                "correcta.",
                    ))
            else:
                filas_docperiodo_ambiguo.append(FilaProblemaMcFinal(
                    fila_excel=numero_fila,
                    documento=mapeada["documento"],
                    ies=mapeada["ies"],
                    tipo="docperiodo_ambiguo",
                    mensaje=f"Ya existe {n_existentes} veces en mc_final (duplicado histórico) — no se tocó, "
                            "revisa manualmente.",
                ))

        filas_aplicadas = 0
        if confirmar and filas_para_aplicar:
            a_insertar = [m for _, m, accion in filas_para_aplicar if accion == "insertar"]
            a_actualizar = [m for _, m, accion in filas_para_aplicar if accion == "actualizar"]

            columnas_sql = ", ".join(MC_FINAL_COLUMNS)
            placeholders = ", ".join(f":{c}" for c in MC_FINAL_COLUMNS)
            stmt_insert = text(f"INSERT INTO analitica_fondos.mc_final ({columnas_sql}) VALUES ({placeholders})")
            for chunk in _chunked(a_insertar):
                if chunk:
                    conn.execute(stmt_insert, chunk)

            if a_actualizar:
                tmp_tabla = "tmp_mc_final_staging"
                conn.execute(text(f"DROP TEMPORARY TABLE IF EXISTS {tmp_tabla}"))
                conn.execute(text(f"CREATE TEMPORARY TABLE {tmp_tabla} LIKE analitica_fondos.mc_final"))
                stmt_insert_tmp = text(f"INSERT INTO {tmp_tabla} ({columnas_sql}) VALUES ({placeholders})")
                for chunk in _chunked(a_actualizar):
                    if chunk:
                        conn.execute(stmt_insert_tmp, chunk)
                set_sql = ", ".join(f"m.{c} = t.{c}" for c in UPDATE_COLUMNS)
                conn.execute(text(
                    f"UPDATE analitica_fondos.mc_final m JOIN {tmp_tabla} t ON m.docperiodo = t.docperiodo "
                    f"SET {set_sql}"
                ))
                conn.execute(text(f"DROP TEMPORARY TABLE IF EXISTS {tmp_tabla}"))

            conn.commit()
            filas_aplicadas = len(filas_para_aplicar)

    problemas = (
        filas_excluidas + filas_invalidas + filas_duplicadas_en_archivo
        + filas_conflicto_ies + filas_docperiodo_ambiguo + filas_renueva_sin_historial
    )
    problemas.sort(key=lambda f: f.fila_excel)

    resumen = ResumenCargaMcFinal(
        confirmado=confirmar,
        periodo=periodo,
        filas_leidas=len(filas_crudas),
        filas_excluidas_manual=len(filas_excluidas),
        filas_invalidas=len(filas_invalidas),
        filas_correctas=len(filas_para_aplicar),
        filas_nuevas=sum(1 for _, _, accion in filas_para_aplicar if accion == "insertar"),
        filas_actualizan=sum(1 for _, _, accion in filas_para_aplicar if accion == "actualizar"),
        filas_duplicadas_en_archivo=len(filas_duplicadas_en_archivo),
        filas_conflicto_ies=len(filas_conflicto_ies),
        filas_docperiodo_ambiguo=len(filas_docperiodo_ambiguo),
        filas_renueva_sin_historial=len(filas_renueva_sin_historial),
        filas_aplicadas=filas_aplicadas,
    )
    return CargaMcFinalResponse(resumen=resumen, problemas=problemas)